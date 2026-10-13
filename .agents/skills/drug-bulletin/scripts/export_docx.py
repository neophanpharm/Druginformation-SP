"""Export Markdown drug bulletin to formatted DOCX with medical typography and strict Vietnamese font support."""
import argparse
import re
import sys
from pathlib import Path

try:
    import docx
    from docx.shared import Inches, Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsdecls, qn
except ImportError:
    print("python-docx is not installed. Please install it with 'pip install python-docx'.", file=sys.stderr)
    sys.exit(1)

FONT_NAME = "Times New Roman"


def apply_font_to_run(run, font_name=FONT_NAME, font_size_pt=None, color_rgb=None, bold=None, italic=None):
    """Ensure font is explicitly set for Latin, Complex Scripts (cs), and EastAsia to eliminate Vietnamese font switching."""
    run.font.name = font_name
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(
        f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}" w:eastAsia="{font_name}"/>'
    )
    # Replace existing rFonts if present
    for existing in rPr.findall(qn('w:rFonts')):
        rPr.remove(existing)
    rPr.append(rFonts)

    if font_size_pt is not None:
        run.font.size = Pt(font_size_pt)
    if color_rgb is not None:
        run.font.color.rgb = color_rgb
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_background(cell, hex_color):
    """Set background fill color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)


def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    """Set cell padding in dxa (1 pt = 20 dxa)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def set_table_borders(table, color="B0C4DE", sz="4", val="single"):
    """Set clean borders for table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def add_formatted_runs(paragraph, text, base_font_size=12, base_bold=False, base_italic=False, base_color=RGBColor(30, 30, 30)):
    """Parse inline markdown tokens (**bold**, *italic*, `code`, <sub>, <sup>) and apply strict fonts."""
    # Pattern to match bold, italic, code, sub, sup
    token_pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|<sub>.*?</sub>|<sup>.*?</sup>)')
    parts = token_pattern.split(text)

    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) >= 4:
            run = paragraph.add_run(part[2:-2])
            apply_font_to_run(run, FONT_NAME, base_font_size, base_color, bold=True, italic=base_italic)
        elif part.startswith('*') and part.endswith('*') and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            apply_font_to_run(run, FONT_NAME, base_font_size, base_color, bold=base_bold, italic=True)
        elif part.startswith('`') and part.endswith('`') and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            apply_font_to_run(run, "Consolas", base_font_size - 1, RGBColor(160, 40, 40), bold=False, italic=False)
        elif part.startswith('<sub>') and part.endswith('</sub>') and len(part) >= 11:
            run = paragraph.add_run(part[5:-6])
            run.font.subscript = True
            apply_font_to_run(run, FONT_NAME, base_font_size - 2, base_color, bold=base_bold, italic=base_italic)
        elif part.startswith('<sup>') and part.endswith('</sup>') and len(part) >= 11:
            run = paragraph.add_run(part[5:-6])
            run.font.superscript = True
            apply_font_to_run(run, FONT_NAME, base_font_size - 2, base_color, bold=base_bold, italic=base_italic)
        else:
            run = paragraph.add_run(part)
            apply_font_to_run(run, FONT_NAME, base_font_size, base_color, bold=base_bold, italic=base_italic)


def markdown_to_docx(md_path, docx_path):
    md_text = Path(md_path).read_text(encoding='utf-8')
    lines = md_text.splitlines()

    doc = docx.Document()

    # Configure A4 page dimensions and medical margins (Left 2.5cm, Right 2cm, Top 2cm, Bottom 2cm)
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.0)

    # Style definitions
    normal_style = doc.styles['Normal']
    normal_style.font.name = FONT_NAME
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(35, 35, 35)

    i = 0
    table_lines = []

    def flush_table(tbl_lines):
        if not tbl_lines:
            return
        rows_data = []
        for l in tbl_lines:
            cells = [c.strip() for c in l.strip().strip('|').split('|')]
            if all(set(c).issubset({'-', ':', ' '}) for c in cells):
                continue
            rows_data.append(cells)

        if not rows_data:
            return

        num_cols = max(len(r) for r in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, color="B0C4DE", sz="4", val="single")

        for row_idx, r_data in enumerate(rows_data):
            row = table.rows[row_idx]
            is_header = (row_idx == 0)
            for col_idx in range(num_cols):
                cell_text = r_data[col_idx] if col_idx < len(r_data) else ""
                cell = row.cells[col_idx]
                set_cell_margins(cell, top=140, bottom=140, left=160, right=160)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15

                # Support line breaks inside table cells like <br>
                cell_lines = cell_text.split('<br>')
                for c_idx, cl in enumerate(cell_lines):
                    if c_idx > 0:
                        p = cell.add_paragraph()
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)
                        p.paragraph_format.line_spacing = 1.15

                    if is_header:
                        set_cell_background(cell, "003366")  # Navy Blue header
                        add_formatted_runs(p, cl.strip(), base_font_size=11, base_bold=True, base_color=RGBColor(255, 255, 255))
                    else:
                        if row_idx % 2 == 1:
                            set_cell_background(cell, "F7F9FB")  # Subtle blue-gray zebra
                        else:
                            set_cell_background(cell, "FFFFFF")
                        add_formatted_runs(p, cl.strip(), base_font_size=10.5, base_color=RGBColor(30, 30, 30))

        # Add clean paragraph spacing after table
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_before = Pt(0)
        spacer.paragraph_format.space_after = Pt(6)

    while i < len(lines):
        line = lines[i]
        raw_strip = line.strip()

        # Check if table line
        if raw_strip.startswith('|') and raw_strip.endswith('|'):
            table_lines.append(raw_strip)
            i += 1
            continue
        elif table_lines:
            flush_table(table_lines)
            table_lines = []

        if not raw_strip:
            i += 1
            continue

        # Horizontal rule divider
        if raw_strip in ('---', '***', '___'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="1" w:color="003366"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            i += 1
            continue

        # Heading 1
        if raw_strip.startswith('# '):
            h = doc.add_paragraph()
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
            add_formatted_runs(h, raw_strip[2:], base_font_size=16, base_bold=True, base_color=RGBColor(0, 51, 102))
            i += 1
            continue

        # Heading 2
        if raw_strip.startswith('## '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            add_formatted_runs(h, raw_strip[3:], base_font_size=13.5, base_bold=True, base_color=RGBColor(0, 60, 120))
            i += 1
            continue

        # Heading 3
        if raw_strip.startswith('### '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(11)
            h.paragraph_format.space_after = Pt(3)
            h.paragraph_format.keep_with_next = True
            add_formatted_runs(h, raw_strip[4:], base_font_size=12.5, base_bold=True, base_color=RGBColor(20, 75, 130))
            i += 1
            continue

        # Heading 4
        if raw_strip.startswith('#### '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(9)
            h.paragraph_format.space_after = Pt(2)
            h.paragraph_format.keep_with_next = True
            add_formatted_runs(h, raw_strip[5:], base_font_size=12, base_bold=True, base_color=RGBColor(40, 40, 40))
            i += 1
            continue

        # Blockquote / Highlight Callout Box
        if raw_strip.startswith('>'):
            quote_text = raw_strip.lstrip('>').strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.right_indent = Inches(0.2)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)

            # Border and shading XML
            pPr = p._p.get_or_add_pPr()
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="12" w:color="003366"/></w:pBdr>')
            p_shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F0F4F8"/>')
            pPr.append(p_border)
            pPr.append(p_shd)

            add_formatted_runs(p, quote_text, base_font_size=11.5, base_color=RGBColor(20, 40, 60))
            i += 1
            continue

        # Bullet List
        if raw_strip.startswith(('* ', '- ', '+ ')):
            bullet_text = raw_strip[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_runs(p, bullet_text, base_font_size=12)
            i += 1
            continue

        # Numbered List
        num_match = re.match(r'^(\d+)\.\s+(.*)$', raw_strip)
        if num_match:
            num_text = num_match.group(2).strip()
            p = doc.add_paragraph(style='List Number')
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_runs(p, num_text, base_font_size=12)
            i += 1
            continue

        # Metadata subtitle (italic)
        if raw_strip.startswith('*') and raw_strip.endswith('*'):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(6)
            add_formatted_runs(p, raw_strip[1:-1], base_font_size=11, base_italic=True, base_color=RGBColor(100, 100, 100))
            i += 1
            continue

        # Regular Body Paragraph
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        add_formatted_runs(p, raw_strip, base_font_size=12)
        i += 1

    if table_lines:
        flush_table(table_lines)

    docx_path = Path(docx_path)
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))
    print(f"Successfully exported DOCX with font {FONT_NAME} to: {docx_path}")


def main():
    parser = argparse.ArgumentParser(description="Export markdown bulletin to formatted DOCX")
    parser.add_argument("--input", required=True, help="Input markdown file path")
    parser.add_argument("--output", required=True, help="Output docx file path")
    args = parser.parse_args()

    markdown_to_docx(args.input, args.output)


if __name__ == '__main__':
    main()
